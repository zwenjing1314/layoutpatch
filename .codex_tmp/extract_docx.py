from docx import Document
p = r'C:\Users\WenJing\Documents\WorkTransfer\layoutpatch\docs\thesis\研究进展与毕业论文实施方案_20261014.docx'
d = Document(p)
print('PARAS', len(d.paragraphs), 'TABLES', len(d.tables))
for i, x in enumerate(d.paragraphs):
    if x.text.strip():
        print(f'P{i}: {x.text}')
print('---TABLES---')
for ti, t in enumerate(d.tables):
    print('TABLE', ti)
    for row in t.rows:
        print(' | '.join(c.text.replace('\n',' / ') for c in row.cells))
