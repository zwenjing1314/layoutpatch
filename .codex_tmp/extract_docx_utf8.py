from docx import Document
p = r'C:\Users\WenJing\Documents\WorkTransfer\layoutpatch\docs\thesis\研究进展与毕业论文实施方案_20261014.docx'
out = r'C:\Users\WenJing\Documents\WorkTransfer\layoutpatch\.codex_tmp\thesis_text_utf8.txt'
d = Document(p)
with open(out, 'w', encoding='utf-8') as f:
    for i, x in enumerate(d.paragraphs):
        if x.text.strip(): f.write(f'P{i}: {x.text}\n')
    f.write('---TABLES---\n')
    for ti, t in enumerate(d.tables):
        f.write(f'TABLE {ti}\n')
        for row in t.rows:
            f.write(' | '.join(c.text.replace('\n',' / ') for c in row.cells)+'\n')
