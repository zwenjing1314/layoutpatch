from pathlib import Path
import re, copy, json, html
from decimal import Decimal, ROUND_HALF_UP
from lxml import etree
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parent.parent
SRC=ROOT/'manuscript-revised.md'
MNS='http://www.w3.org/1998/Math/MathML'
XSL=etree.XSLT(etree.parse(r'C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL'))
SYMS={'nu':'ν','ell':'ℓ','top':'⊤','in':'∈','mathbb':'','times':'×','le':'≤','ge':'≥','neq':'≠','varnothing':'∅','infty':'∞','rho':'ρ','mu':'μ','phi':'φ','omega':'ω','alpha':'α','delta':'δ','sigma':'σ','tau':'τ','epsilon':'ϵ','varepsilon':'ε','beta':'β','sum':'∑','cup':'∪','cap':'∩','ldots':'…','rightarrow':'→','bar':'','tilde':'','hat':'','max':'max','log':'log','exp':'exp','sin':'sin','cos':'cos'}
def mel(tag,*kids,text=None,**attrs):
    e=etree.Element('{'+MNS+'}'+tag,**attrs)
    if text is not None:e.text=text
    for k in kids:
        if k is not None:e.append(k)
    return e
class MathParser:
    def __init__(self,s):self.s=s;self.i=0
    def seq(self,stop=None):
        kids=[]
        while self.i<len(self.s):
            if stop and self.s[self.i]==stop:self.i+=1;break
            if self.s[self.i].isspace():self.i+=1;continue
            a=self.atom()
            if a is None:continue
            sub=sup=None
            while self.i<len(self.s) and self.s[self.i] in "_^'":
                op=self.s[self.i];self.i+=1
                if op=="'":v=mel('mo',text='′');sup=mel('mrow',sup,v) if sup is not None else v
                else:
                    v=self.atom()
                    if op=='_':sub=v
                    else:sup=v
            if sub is not None and sup is not None:a=mel('msubsup',a,sub,sup)
            elif sub is not None:a=mel('msub',a,sub)
            elif sup is not None:a=mel('msup',a,sup)
            kids.append(a)
        return mel('mrow',*kids)
    def group(self):
        while self.i<len(self.s) and self.s[self.i].isspace():self.i+=1
        if self.i<len(self.s) and self.s[self.i]=='{':self.i+=1;return self.seq('}')
        return self.atom()
    def rawgroup(self):
        a=self.group();return ''.join(a.itertext())
    def atom(self):
        while self.i<len(self.s) and self.s[self.i].isspace():self.i+=1
        if self.i>=len(self.s):return mel('mrow')
        c=self.s[self.i];self.i+=1
        if c=='{':return self.seq('}')
        if c=='\\':
            m=re.match(r'[A-Za-z]+',self.s[self.i:])
            if not m:
                c=self.s[self.i] if self.i<len(self.s) else '';self.i+=1
                return mel('mspace',width='.15em') if c in ',;! ' else mel('mo',text='‖' if c=='|' else c)
            cmd=m.group();self.i+=len(cmd)
            if cmd in ('left','right','big','Big','bigl','bigr'):return self.atom()
            if cmd in ('quad','qquad'):return mel('mspace',width='1em')
            if cmd=='frac':return mel('mfrac',self.group(),self.group())
            if cmd=='sqrt':return mel('msqrt',self.group())
            if cmd in ('bar','tilde','hat'):
                return mel('mover',self.group(),mel('mo',text={'bar':'¯','tilde':'˜','hat':'^'}[cmd]),accent='true')
            if cmd in ('mathrm','operatorname','text'):return mel('mi',text=self.rawgroup(),mathvariant='normal')
            if cmd=='mathbb':return mel('mi',text=self.rawgroup(),mathvariant='double-struck')
            if cmd=='mathcal':return mel('mi',text=self.rawgroup(),mathvariant='script')
            if cmd not in SYMS:raise ValueError('Unhandled math command '+cmd)
            val=SYMS[cmd];return mel('mo' if cmd in ('cup','cap','in','le','ge','neq','times','rightarrow','top') else 'mi',text=val)
        if c.isdigit():
            m=re.match(r'[0-9.]*',self.s[self.i:]);v=c+m.group();self.i+=len(m.group());return mel('mn',text=v)
        if c.isalpha():return mel('mi',text=c)
        return mel('mo',text=c)
def math_node(s):
    if '\\begin{cases}' in s:
        before,rest=s.split('\\begin{cases}',1);inside,after=rest.split('\\end{cases}',1)
        rows=[]
        for row in inside.split('\\\\'):
            rows.append(mel('mtr',*[mel('mtd',MathParser(c).seq()) for c in row.split('&')]))
        tree=mel('math',MathParser(before).seq(),mel('mfenced',mel('mtable',*rows),open='{',close=''),MathParser(after).seq())
    else:tree=mel('math',MathParser(s).seq())
    out=XSL(tree).getroot()
    om=out if out.tag==qn('m:oMath') else out.find('.//'+qn('m:oMath'))
    if om is None:raise ValueError('No OMML '+s)
    return copy.deepcopy(om)

text=SRC.read_text(encoding='utf-8-sig')
# Move the lengthy symbol table to the appendix without dropping its content.
sym=re.search(r'\*\*表3-1 主要符号与维度\*\*\s*\n(\|.*?)(?=\n## 3\.2)',text,re.S)
symbol_table=sym.group(1).strip()
text=text[:sym.start()]+'主要符号及维度见附录A表7。\n\n'+text[sym.end():]
text=re.sub(r'\*\*图3-1 方法流程图草案\*\*.*?各方法分别输出排序，图末端不表示分数集成。',
    '各方法的平行关系及D的局部融合结构见图2。C与D分别从原始表示出发，E复用B的分数作为基值；标注证据只进入训练监督支路。\n\n![图2](figures/figure-2-methods.png)\n\n图2 方法结构与局部融合。E为事后补充，C/D为平行分支；虚线区域为训练监督。',text,flags=re.S)
text=text.replace('## 3.7 方法差异表与流程图草案','## 3.7 方法比较')
text=text.replace('参数量沿用第1轮材料的登记值，网络维度已与本轮实现核对；本轮未运行参数计数。','表1列出各模块的登记可训练参数量。')
text=text.replace('**表3-2 方法差异**','**表1 方法差异**')
text=text.replace('**表5-1','**表3').replace('**表5-2','**表4').replace('**表5-3','**表5').replace('**表5-4','**表6')
text=text.replace('表3-1','表7').replace('表3-2','表1').replace('表5-1','表3').replace('表5-2','表4').replace('表5-3','表5').replace('表5-4','表6')
text=text.replace('式（15）先对原始','式（16）先对原始').replace('式（19）直接优化 patch','式（20）直接优化 patch')
text=text.replace('退化节点返回全零行','退化节点的空间对应权重为零')
text=text.replace('空 (W_{\\mathrm{raw}}) 候选',r'空 \(W_{\mathrm{raw}}\) 候选')
text=text.replace('本轮没有进行替换汇聚函数、改变读出层或逐层归因的机制实验','本文尚未通过机制实验检验这些假设')
text=text.replace('并未在本轮逐项复现','原版本未完整追溯')
text=text.replace('E 的候选级训练单元、模型结构和 checkpoint 选择分母仍与 C/D 不同。','E 的候选级训练单元和模型结构与 C/D 不同；checkpoint 选择均使用 source_val 的374题、448个评测单元。')
text=text.replace('；不加载 smoke、C 或旧融合权重','，各方法不加载其他适配器或融合器的既有权重')
# Figure 1 follows the task definition.
anchor='冻结 ColPali 的视觉与问题编码部分'
pos=text.index(anchor)
text=text[:pos]+'任务与空间对应关系见图1。\n\n![图1](figures/figure-1-task.png)\n\n图1 任务与空间对应示意。示意图，非实验样本；蓝框仅表示解析候选，不代表标注证据或成功预测。\n\n'+text[pos:]
data='''**表2 数据角色与实际分母**

| 来源与角色 | 用途或口径 | 文档数 | 题数 | 唯一页数 | 题—页单元数 |
|---|---|---:|---:|---:|---:|
| MMDocIR source_train | 原始角色 / 区域审计 | 169 | 826 | — | 1059 |
| C/D训练子集 | patch级损失可用 | — | — | — | 1002 |
| E训练子集 | 候选级进一步筛选 | — | — | — | 843 |
| MMDocIR source_val | 原始角色 / 区域审计 | 43 | 375 | — | 450 |
| source_val有效评测 | 权重选择 | — | 374 | — | 448 |
| MMDocIR source_dev | 开发比较 | 91 | 412 | — | 545 |
| quarantine | 隔离，不参与比较 | 10 | 45 | — | 53 |
| SciEGQA固定test | 原冻结外部与E补充 | — | 1623 | 837 | 1941 |

注：题数、唯一页数与题—页单元数为不同计数口径；“—”表示正文未提供该字段，不代表零。外部全部方法共享9746个候选。MMDocIR官方evaluation在本文重新划为训练、验证和开发来源，非最终独立测试。

'''
pos=text.index('## 4.3 候选、监督与过滤')
text=text[:pos]+'数据角色与实际分母汇总见表2。\n\n'+data+text[pos:]
# Precision changes only round pre-existing exact values.
body,refs=text.split('# 参考文献',1)
body=re.sub(r'(?<![\w])([+−-]?\d+\.\d{5,})(?!\d)',lambda m:format(Decimal(m.group(1).replace('−','-')).quantize(Decimal('.0001'),rounding=ROUND_HALF_UP),'f').replace('-', '−') if m.group(1).startswith('−') else (('+' if m.group(1).startswith('+') else '')+format(Decimal(m.group(1)).quantize(Decimal('.0001'),rounding=ROUND_HALF_UP),'f')),body)
text=body+'# 参考文献'+refs
text=re.sub(r'\(\(([\s\S]*?)\)\)',r'(\1)',text) if False else text
# Plain parentheses around numeric intervals are prose, not equation environments.
text=re.sub(r'\(([+−-]\d+\.\d+)\)',r'\1',text)
text=re.sub(r'\((\[[+−\-\d., ]+\])\)',r'\1',text)
text=text.replace('\n---\n','\n')
text+='\n\n# 附录A 主要符号与维度\n\n**表7 主要符号与维度**\n\n'+symbol_table+'\n'
(ROOT/'manuscript-illustrated.md').write_text(text,encoding='utf-8')

doc=Document()
sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.left_margin=sec.right_margin=Inches(.8);sec.top_margin=Inches(.75);sec.bottom_margin=Inches(.72)
sec.header_distance=sec.footer_distance=Inches(.3)
styles=doc.styles
for style in styles:
    for border in list(style.element.findall('.//'+qn('w:pBdr'))):border.getparent().remove(border)
for name in ['Normal','Title','Heading 1','Heading 2','Heading 3','Caption']:
    s=styles[name];s.font.name='Times New Roman';s.font.color.rgb=RGBColor(0,0,0)
    s._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'宋体' if name in ['Normal','Caption'] else '黑体')
styles['Normal'].font.size=Pt(10.5)
styles['Normal'].paragraph_format.line_spacing=1.18
styles['Normal'].paragraph_format.space_after=Pt(4)
styles['Normal'].paragraph_format.widow_control=True
styles['Title'].font.size=Pt(19);styles['Title'].paragraph_format.space_after=Pt(14)
for name,size in [('Heading 1',14),('Heading 2',12),('Heading 3',11)]:
    styles[name].font.size=Pt(size);styles[name].font.bold=True
    styles[name].paragraph_format.space_before=Pt(12);styles[name].paragraph_format.space_after=Pt(6)
    styles[name].paragraph_format.keep_with_next=True
styles['Caption'].font.size=Pt(9);styles['Caption'].paragraph_format.space_after=Pt(6)
doc.core_properties.author='';doc.core_properties.last_modified_by=''
doc.core_properties.title='布局条件适配与直接排序的实证比较'
settings=doc.settings.element
mp=OxmlElement('m:mathPr');mf=OxmlElement('m:mathFont');mf.set(qn('m:val'),'Cambria Math');mp.append(mf);settings.append(mp)
foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)

def add_inline(p,s,bold=False):
    # Markdown emphasis/code and both forms of math delimiters.
    pat=r'(\\\([\s\S]*?\\\)|\$[^$]+\$|\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)'
    for part in re.split(pat,s):
        if not part:continue
        if part.startswith('\\('):p._p.append(math_node(part[2:-2]))
        elif part.startswith('$'):p._p.append(math_node(part[1:-1]))
        else:
            b=bold or part.startswith('**');it=part.startswith('*') and not b
            val=part.strip('*') if part.startswith('*') else part.strip('`') if part.startswith('`') else part
            r=p.add_run(val);r.bold=b;r.italic=it
            r.font.name='Times New Roman';r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'黑体' if p.style.name=='Title' or p.style.name.startswith('Heading') else '宋体')

def add_table(lines):
    rows=[[c.strip() for c in l.strip().strip('|').split('|')] for l in lines]
    rows=[r for r in rows if not all(re.fullmatch(r'[:\-]+',c) for c in r)]
    n=len(rows[0]);table=doc.add_table(rows=0,cols=n);table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
    widths={3:[1.35,1.45,4.1],4:[1.7,.8,2.2,2.2],5:[2.15,1.18,1.18,1.08,1.31],6:[.67,1.79,1.25,1.56,.65,.98]}[n]
    if rows[0][0]=='来源与角色':widths=[1.69,1.7,.68,.68,.75,1.0]
    for col,w in zip(table.columns,widths):col.width=Inches(w)
    for ri,row in enumerate(rows):
        cells=table.add_row().cells
        for ci,(cell,value) in enumerate(zip(cells,row)):
            cell.width=Inches(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc=cell._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                b=OxmlElement('w:'+side);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
            tc.append(borders)
            mar=OxmlElement('w:tcMar')
            for side,v in [('top','65'),('bottom','65'),('left','80'),('right','80')]:
                x=OxmlElement('w:'+side);x.set(qn('w:w'),v);x.set(qn('w:type'),'dxa');mar.append(x)
            tc.append(mar)
            if ri==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EAF0F5');tc.append(sh)
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.05
            p.paragraph_format.keep_with_next=ri<len(rows)-1
            p.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 or n==6 and rows[0][0]=='方法' else WD_ALIGN_PARAGRAPH.CENTER
            add_inline(p,value,bold=ri==0)
            for r in p.runs:r.font.size=Pt(8.5 if n==6 else 9)
        pr=table.rows[-1]._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        if ri==0:pr.append(OxmlElement('w:tblHeader'))
    doc.add_paragraph().paragraph_format.space_after=Pt(1)

lines=text.splitlines();i=0;first=True;figure_count=0;eqcount=0
while i<len(lines):
    line=lines[i].strip()
    if not line:i+=1;continue
    if line.startswith('#'):
        level=len(line)-len(line.lstrip('#'));s=line[level:].strip()
        if first:
            p=doc.add_paragraph(style='Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            # Readability adjustment: title in two balanced lines.
            add_inline(p,s.replace('：','\n'));first=False
        else:
            p=doc.add_paragraph(s,style=f'Heading {min(level,3)}')
            if s in ['摘要','参考文献'] or s.startswith('附录A'):p.paragraph_format.page_break_before=s!='摘要'
        i+=1;continue
    if line=='\\[':
        i+=1;eq=[]
        while lines[i].strip()!='\\]':eq.append(lines[i]);i+=1
        eq='\n'.join(eq);num=re.search(r'\\tag\{(\d+)\}',eq).group(1);eq=re.sub(r'\\tag\{\d+\}','',eq).strip()
        p=doc.add_paragraph();p.paragraph_format.space_before=Pt(5);p.paragraph_format.space_after=Pt(5)
        p.paragraph_format.keep_together=True;p.paragraph_format.line_spacing=1.05
        # Long equations are broken at an existing semantic separator.
        bits=eq.split('\\qquad') if num in ['6','12','14','20','21'] else [eq]
        p.paragraph_format.tab_stops.add_tab_stop(Inches(3.3),WD_TAB_ALIGNMENT.CENTER)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(6.85),WD_TAB_ALIGNMENT.RIGHT)
        for k,bit in enumerate(bits):
            if k:p.add_run().add_break()
            p.add_run('\t');p._p.append(math_node(bit))
        p.add_run('\t('+num+')');eqcount+=1;i+=1;continue
    if line.startswith('|'):
        tl=[]
        while i<len(lines) and lines[i].strip().startswith('|'):tl.append(lines[i]);i+=1
        add_table(tl);continue
    if line.startswith('!['):
        fn=re.search(r'\(([^)]+)\)',line).group(1)
        p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next=True;p.paragraph_format.keep_together=True
        p.add_run().add_picture(str(ROOT/fn),width=Inches(6.75));figure_count+=1;i+=1;continue
    if line.startswith('**表'):
        p=doc.add_paragraph(style='Caption');p.paragraph_format.keep_with_next=True;p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        add_inline(p,line);i+=1;continue
    if re.match(r'^图[12] ',line):
        p=doc.add_paragraph(style='Caption');p.paragraph_format.keep_with_next=False;add_inline(p,line);i+=1;continue
    p=doc.add_paragraph()
    p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    if not line.startswith(('注：','[','**关键词')):p.paragraph_format.first_line_indent=Pt(21)
    if line.startswith('['):p.paragraph_format.first_line_indent=Pt(-18);p.paragraph_format.left_indent=Pt(18)
    add_inline(p,line);i+=1
doc.save(ROOT/'paper-advisor-v1.docx')
(ROOT/'figures'/'build-summary.json').write_text(json.dumps({'figures':figure_count,'tables':len(doc.tables),'display_equations':eqcount,'paragraphs':len(doc.paragraphs)},ensure_ascii=False,indent=2),encoding='utf-8')
print('DOCX written; figures',figure_count,'tables',len(doc.tables),'equations',eqcount)
